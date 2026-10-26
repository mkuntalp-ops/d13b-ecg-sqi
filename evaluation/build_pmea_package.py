"""
Build the Physiological Measurement (IOP) submission package from the assembled manuscript (2026-10-26).

Inputs: manuscript/main_q1.md and manuscript/supplement_q1.md (from evaluation/assemble_q1_manuscript.py),
manuscript/tripod_ai_checklist_q1.md, manuscript/PMEA_cover_letter_draft.md, figures/q1_figure{1,2,3}.pdf.

Transformations for the main text (no numbers are changed):
- author block after the title (as in the author's previous PMEA submission PMEA-107866);
- abstract headings in the PMEA form (Objective / Approach / Main results / Significance), keywords and abbreviations;
- numbered citations [n] converted to IOP Harvard author-year citations, ordered by year within a parenthesis; only
  integer-only brackets in the main text are converted (all were checked to be citations); the reference list is
  rewritten in IOP Harvard style, sorted alphabetically;
- the declaration sections are replaced by the PMEA set (Acknowledgements, AI disclosure statement, Funding, Ethics
  approval and consent, Conflicts of interest, Author contributions, Data availability statement). Statements that
  the author must confirm are marked [AUTHOR TO CONFIRM].
Word files are produced with pandoc (A4 portrait reference document) and post-processed with python-docx: table
borders, small table fonts, repeated header rows, rows that do not split, captions kept with their tables.

Output: PMEA_submission_v1/ (Word files, vector figures, Markdown sources, manifest with SHA-256 hashes)
Usage: python evaluation/build_pmea_package.py
"""

import hashlib
import json
import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
MS = ROOT / "manuscript"
OUT = ROOT / "PMEA_submission_v1"
SRC = OUT / "markdown_sources"
V = "v1"

AUTHOR_BLOCK = """Mehmet Kuntalp

Department of Electrical and Electronics Engineering, Dokuz Eylül University, İzmir, Turkey. Postal address: Tınaztepe Campus, Buca, 35160 İzmir, Turkey

Corresponding author: Mehmet Kuntalp, mehmet.kuntalp@deu.edu.tr

ORCID: Mehmet Kuntalp, https://orcid.org/0000-0002-3381-9026
"""

ABBREVIATIONS = ("**Abbreviations:** AP, average precision; AUROC, area under the receiver operating characteristic curve; "
                 "CI, confidence interval; CNN, convolutional neural network; FPR, false-positive rate; FR, false reject rate; "
                 "LODO, leave-one-dataset-out; MLP, multilayer perceptron; NPV, negative predictive value; PPV, positive "
                 "predictive value; SNR, signal-to-noise ratio; SQI, signal quality index.")

# number -> (in-text citation, year, IOP Harvard entry)
REFS = {
    1: ("Satija et al 2018", 2018, "Satija U, Ramkumar B and Manikandan M S 2018 A review of signal processing techniques for electrocardiogram signal quality assessment *IEEE Reviews in Biomedical Engineering* 11 36–52 https://doi.org/10.1109/RBME.2018.2810957"),
    2: ("Clifford et al 2012", 2012, "Clifford G D, Behar J, Li Q and Rezek I 2012 Signal quality indices and data fusion for determining clinical acceptability of electrocardiograms *Physiological Measurement* 33 1419–1433 https://doi.org/10.1088/0967-3334/33/9/1419"),
    3: ("Orphanidou et al 2015", 2015, "Orphanidou C, Bonnici T, Charlton P, Clifton D, Vallance D and Tarassenko L 2015 Signal-quality indices for the electrocardiogram and photoplethysmogram: derivation and applications to wireless monitoring *IEEE Journal of Biomedical and Health Informatics* 19 832–838 https://doi.org/10.1109/JBHI.2014.2338351"),
    4: ("Li et al 2008", 2008, "Li Q, Mark R G and Clifford G D 2008 Robust heart rate estimation from multiple asynchronous noisy sources using signal quality indices and a Kalman filter *Physiological Measurement* 29 15–32 https://doi.org/10.1088/0967-3334/29/1/002"),
    5: ("Liu et al 2022", 2022, "Liu F et al 2022 Wearable electrocardiogram quality assessment using wavelet scattering and LSTM *Frontiers in Physiology* 13 905447 https://doi.org/10.3389/fphys.2022.905447"),
    6: ("Hu et al 2025", 2025, "Hu Y, Liu Q, Zhou Z, Xu W and Tang H 2025 A paralleled multi-task learning-based framework for single-lead ECG fine-grained noise localization, denoising and signal quality assessment *Sensors* 25 7152 https://doi.org/10.3390/s25237152"),
    7: ("Huang et al 2024", 2024, "Huang X, Zhang F, Fan H, Chang H, Zhou B and Li Z 2024 Pseudo anomalies enhanced deep support vector data description for electrocardiogram quality assessment *Computers in Biology and Medicine* 170 107928 https://doi.org/10.1016/j.compbiomed.2024.107928"),
    8: ("Herraiz et al 2020", 2020, "Herraiz Á H, Martínez-Rodrigo A, Bertomeu-González V, Quesada A, Rieta J J and Alcaraz R 2020 A deep learning approach for featureless robust quality assessment of intermittent atrial fibrillation recordings from portable and wearable devices *Entropy* 22 733 https://doi.org/10.3390/e22070733"),
    9: ("Liquid AI 2026", 2026, "Liquid AI 2026 LiquidAI/d1-3B: open-weight edge multimodal decision model (Hugging Face model repository, revision 051bcc464b01b9f92942b364d9586b0ef5912432) https://huggingface.co/LiquidAI/d1-3B"),
    10: ("Smital et al 2020", 2020, "Smital L et al 2020 Real-time quality assessment of long-term ECG signals recorded by wearables in free-living conditions *IEEE Transactions on Biomedical Engineering* 67 2721–2734 https://doi.org/10.1109/TBME.2020.2969719"),
    11: ("Silva et al 2011", 2011, "Silva I, Moody G B and Celi L 2011 Improving the quality of ECGs collected using mobile phones: the PhysioNet/Computing in Cardiology Challenge 2011 *Computing in Cardiology* 38 273–276"),
    12: ("PhysioNet 2011", 2011, "PhysioNet 2011 Improving the quality of ECGs collected using mobile phones: the PhysioNet/Computing in Cardiology Challenge 2011 (challenge description) https://moody-challenge.physionet.org/2011/"),
    13: ("Subbaswamy and Saria 2020", 2020, "Subbaswamy A and Saria S 2020 From development to deployment: dataset shift, causality, and shift-stable models in health AI *Biostatistics* 21 345–352 https://doi.org/10.1093/biostatistics/kxz041"),
    14: ("Guo et al 2022", 2022, "Guo L L et al 2022 Evaluation of domain generalization and adaptation on improving model robustness to temporal dataset shift in clinical medicine *Scientific Reports* 12 2726 https://doi.org/10.1038/s41598-022-06484-1"),
    15: ("Moeyersons et al 2019", 2019, "Moeyersons J et al 2019 Artefact detection and quality assessment of ambulatory ECG signals *Computer Methods and Programs in Biomedicine* 182 105050 https://doi.org/10.1016/j.cmpb.2019.105050"),
    16: ("Hollmann et al 2025", 2025, "Hollmann N et al 2025 Accurate predictions on small data with a tabular foundation model *Nature* 637 319–326 https://doi.org/10.1038/s41586-024-08328-6"),
    17: ("Coppola et al 2024", 2024, "Coppola E, Savardi M, Massussi M, Adamo M, Metra M and Signoroni A 2024 HuBERT-ECG as a self-supervised foundation model for broad and scalable cardiac applications *medRxiv* (preprint) https://doi.org/10.1101/2024.11.14.24317328"),
    18: ("McKeen et al 2025", 2025, "McKeen K, Masood S, Toma A, Rubin B and Wang B 2025 ECG-FM: an open electrocardiogram foundation model *JAMIA Open* 8 ooaf122 https://doi.org/10.1093/jamiaopen/ooaf122"),
    19: ("Nemcova et al 2020", 2020, "Nemcova A, Smisek R, Opravilová K, Vitek M, Smital L and Maršánová L 2020 Brno University of Technology ECG Quality Database (BUT QDB) (version 1.0.0) *PhysioNet* https://doi.org/10.13026/kah4-0w24"),
    20: ("Smital et al 2026", 2026, "Smital L et al 2026 Real-world 24h+ ECG dataset with quality annotations and motion context *Scientific Data* https://doi.org/10.1038/s41597-026-07905-w"),
    21: ("Clifford et al 2017", 2017, "Clifford G D et al 2017 AF classification from a short single lead ECG recording: the PhysioNet/Computing in Cardiology Challenge 2017 *Computing in Cardiology* 44 https://doi.org/10.22489/CinC.2017.065-469"),
    22: ("Moody et al 1984", 1984, "Moody G B, Muldrow W E and Mark R G 1984 A noise stress test for arrhythmia detectors *Computers in Cardiology* 11 381–384"),
    23: ("Pan and Tompkins 1985", 1985, "Pan J and Tompkins W J 1985 A real-time QRS detection algorithm *IEEE Transactions on Biomedical Engineering* BME-32 230–236 https://doi.org/10.1109/TBME.1985.325532"),
    24: ("DeLong et al 1988", 1988, "DeLong E R, DeLong D M and Clarke-Pearson D L 1988 Comparing the areas under two or more correlated receiver operating characteristic curves: a nonparametric approach *Biometrics* 44 837–845 https://doi.org/10.2307/2531595"),
    25: ("Holm 1979", 1979, "Holm S 1979 A simple sequentially rejective multiple test procedure *Scandinavian Journal of Statistics* 6 65–70"),
    26: ("Grinsztajn et al 2026", 2026, "Grinsztajn L et al 2026 TabPFN-3: technical report *arXiv* 2605.13986 https://arxiv.org/abs/2605.13986"),
}

DECLARATIONS = """## Acknowledgements

The author thanks the investigators and institutions that created and made available the BUT QDB, the PhysioNet/Computing in Cardiology Challenge 2011 and 2017 datasets and the MIT-BIH Noise Stress Test Database, and the developers of the openly released models used in this study.

## AI disclosure statement

[AUTHOR TO CONFIRM] Generative-AI tools were used in three distinct roles.

- *Antigravity (Google):* used to write an initial version of the analysis code and of the manuscript draft.
- *Claude Code (Anthropic):* used to audit and correct that code, including independent recomputation of the reported metrics and re-execution checks; to write the later analysis, table-generation, figure-generation and verification scripts; and to revise and edit the manuscript text.
- *Perplexity:* used as an LLM-as-a-judge. It reviewed successive versions of the manuscript, Supplementary Material and reporting checklist and commented on methodological, statistical and reporting problems, numerical consistency and reference accuracy.

Each critique was checked against the stored study outputs before any change was accepted, and suggestions not supported by the data were not adopted.

The author defined the research question, scientific objectives, datasets, evaluation protocol and interpretation framework; executed all code locally; reviewed the implementation, analytic outputs, reported results and references; and retains full responsibility for the study design, data handling, methods, analyses, results, interpretations, originality and integrity of this manuscript.

## Funding

[AUTHOR TO CONFIRM] This research received no specific grant from any funding agency in the public, commercial or not-for-profit sectors.

## Ethics approval and consent

[AUTHOR TO CONFIRM] This study is a secondary analysis of publicly available, de-identified ECG datasets, used under the terms of their public releases: BUT QDB (collected under IRB protocol EC:EK:05b/2018 of Brno University of Technology), the PhysioNet/Computing in Cardiology Challenge 2011 and 2017 datasets, and the MIT-BIH Noise Stress Test Database. No new data were collected from human participants, and the author had no access to identifiable information. Ethical approval was not required for this secondary analysis of public de-identified data.

## Conflicts of interest

[AUTHOR TO CONFIRM] The author declares no competing interests.

## Author contributions

[AUTHOR TO CONFIRM] Mehmet Kuntalp (sole author): Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing – original draft, Writing – review & editing.

## Data availability statement

The source ECG data are publicly available from PhysioNet: BUT QDB (Nemcova et al 2020), the PhysioNet/Computing in Cardiology Challenge 2011 (Silva et al 2011) and 2017 (Clifford et al 2017) datasets, and the MIT-BIH Noise Stress Test Database (Moody et al 1984).

The repository at https://github.com/mkuntalp-ops/d13b-ecg-sqi (MIT licence) contains:
- the analysis code;
- the dated analysis plan;
- the cohort and fold manifests;
- the saved predictions of every model;
- the scripts that regenerate every table and figure from those predictions.

The following are not redistributed:
- raw ECG waveforms;
- the weights of LiquidAI/d1-3B, HuBERT-ECG and TabPFN, which are available from their providers under their own licences (HuBERT-ECG: CC BY-NC 4.0).

[AUTHOR TO ADD: archive DOI of the release, if archived.]
"""


def cite(nums):
    items = sorted({n for n in nums}, key=lambda n: (REFS[n][1], REFS[n][0]))
    return "(" + ", ".join(REFS[n][0] for n in items) + ")"


def expand(spec):
    nums = []
    for part in re.split(r"\s*,\s*", spec):
        if "–" in part:
            a, b = (int(x) for x in part.split("–"))
            nums += list(range(a, b + 1))
        else:
            nums.append(int(part))
    return nums


def convert_main(text: str) -> str:
    title_end = text.index("\n", text.index("# "))
    text = text[:title_end + 1] + "\n" + AUTHOR_BLOCK + text[title_end + 1:]
    for a, b in [("**Methods.**", "**Approach.**"), ("**Results.**", "**Main results.**"), ("**Conclusions.**", "**Significance.**")]:
        assert text.count(a) == 1, a
        text = text.replace(a, b)
    kw = text.index("**Keywords:**")
    kw_end = text.index("\n", kw)
    text = text[:kw_end + 1] + "\n" + ABBREVIATIONS + "\n" + text[kw_end + 1:]
    body, _ = text.split("## References", 1)
    start = body.index("## Ethics, funding and competing interests")
    body = body[:start] + DECLARATIONS + "\n"
    pat = re.compile(r"\[(\d{1,2}(?:\s*[–,]\s*\d{1,2})*)\]")
    found = [expand(m.group(1)) for m in pat.finditer(body)]
    assert all(1 <= n <= 26 for f in found for n in f)
    assert {n for f in found for n in f} == set(REFS), "every reference must be cited"
    body = pat.sub(lambda m: cite(expand(m.group(1))), body)
    refs = sorted(REFS.values(), key=lambda r: (r[2].split(" ")[0].lower(), r[1], r[2].lower()))  # first author, year
    body += "## References\n\n" + "\n\n".join(r[2] for r in refs) + "\n"
    return body.replace("../figures/", str(ROOT / "figures") + "/")


def reference_docx(path: Path):
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(str(path))
    for s in d.sections:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.2)
    for name in ["Normal", "Body Text", "First Paragraph", "Compact"]:
        if name in [st.name for st in d.styles]:
            d.styles[name].font.name = "Times New Roman"
            d.styles[name].font.size = Pt(11)
    d.save(str(path))


def _set_cell_borders(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "999999")
        borders.append(e)
    tbl.tblPr.append(borders)


def postprocess(path: Path, base_pt=8.0, wide_pt=7.0):
    d = Document(str(path))
    body = d.element.body
    for table in d.tables:
        ncol = len(table.columns)
        size = Pt(wide_pt if ncol >= 7 else base_pt)
        _set_cell_borders(table)
        for i, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                h = OxmlElement("w:tblHeader")
                trPr.append(h)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = size
                    if i < len(table.rows) - 1:
                        p.paragraph_format.keep_with_next = True
        # keep the caption block (paragraphs just before the table) with the table
        prev = table._tbl.getprevious()
        for _ in range(6):
            if prev is None or prev.tag != qn("w:p"):
                break
            pPr = prev.get_or_add_pPr()
            kn = OxmlElement("w:keepNext")
            pPr.append(kn)
            text = "".join(t.text or "" for t in prev.iter(qn("w:t")))
            if text.startswith("Table "):
                break
            prev = prev.getprevious()
    for p in d.paragraphs:  # keep figures with their captions
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
    d.save(str(path))


def to_docx(md: Path, out: Path, ref: Path, **kw):
    subprocess.run(["pandoc", str(md), "--from", "markdown+pipe_tables+lists_without_preceding_blankline", "--reference-doc", str(ref),
                    "--resource-path", f"{MS}:{ROOT}", "-o", str(out)], check=True)
    postprocess(out, **kw)


def title_page(title: str) -> str:
    return f"""# {title}

Physiological Measurement — Research paper

**Author:** Mehmet Kuntalp (sole author) [AUTHOR TO CONFIRM]

**Affiliation:** Department of Electrical and Electronics Engineering, Dokuz Eylül University, İzmir, Turkey

**Corresponding author:** Mehmet Kuntalp, Department of Electrical and Electronics Engineering, Dokuz Eylül University, İzmir, Turkey; mehmet.kuntalp@deu.edu.tr; postal address: Tınaztepe Campus, Buca, 35160 İzmir, Turkey

**ORCID:** 0000-0002-3381-9026
"""


def declarations_form(main_md: str) -> str:
    sec = lambda name: main_md.split(f"## {name}\n", 1)[1].split("\n## ", 1)[0].strip()
    items = [("Author", "Mehmet Kuntalp (sole author) [AUTHOR TO CONFIRM]. Affiliation: Department of Electrical and Electronics Engineering, Dokuz Eylül University, İzmir, Turkey. Corresponding author: Mehmet Kuntalp. Email: mehmet.kuntalp@deu.edu.tr. Postal address: Tınaztepe Campus, Buca, 35160 İzmir, Turkey. ORCID: 0000-0002-3381-9026."),
             ("Funding and sponsor roles", sec("Funding")), ("Financial/non-financial conflicts", sec("Conflicts of interest")),
             ("Ethics and consent", sec("Ethics approval and consent")),
             ("Author contributions", sec("Author contributions") + " Acknowledgements: " + sec("Acknowledgements")),
             ("Data availability", sec("Data availability statement")), ("AI disclosure", sec("AI disclosure statement")),
             ("Related submission", "A separate manuscript by the same author, PMEA-107866 (\"Teacher–Student Distillation of TabPFN for Retrospective Two-Lead ECG Features\"), is under consideration at Physiological Measurement. The two manuscripts share no datasets, analyses or results; both use TabPFN, in different roles (see cover letter)."),
             ("Reporting guideline", "TRIPOD+AI; the completed checklist is provided as a supplementary file.")]
    return "# Author information and declarations\n\n" + "\n\n".join(f"**{i}. {k}.** {v}" for i, (k, v) in enumerate(items, 1)) + "\n"


CHROME = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")


def previews():
    """Preview PDFs from HTML (headless Chrome); not for the portal - the portal PDFs are exported from Word."""
    if not CHROME.exists():
        print("[preview] Chrome not found; previews skipped")
        return
    (OUT / "PDF_preview").mkdir(exist_ok=True)
    pat = re.compile(r'(<p><em>Table \d+\.(?:(?!</p>).)*</p>\s*(?:<ul>(?:(?!</ul>).)*</ul>\s*)?<table(?:(?!</table>).)*</table>)', re.S)
    for name in (f"PMEA_Manuscript_{V}", f"PMEA_Supplementary_Material_1_{V}"):
        html = SRC / f"{name}.html"
        subprocess.run(["pandoc", str(SRC / f"{name}.md"), "-f", "markdown+lists_without_preceding_blankline", "-s",
                        "--embed-resources", "--resource-path", f"{MS}:{ROOT}", "--css", str(ROOT / "evaluation" / "pmea_print.css"),
                        "--metadata", "title= ", "-o", str(html)], check=True, stderr=subprocess.DEVNULL)
        h = html.read_text(encoding="utf-8")
        html.write_text(pat.sub(lambda m: f'<div class="tblock">{m.group(1)}</div>', h), encoding="utf-8")
        subprocess.run([str(CHROME), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={OUT / 'PDF_preview' / (name + '_preview.pdf')}", html.as_uri()],
                       check=True, stderr=subprocess.DEVNULL)
        html.unlink()


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "PMEA_Main_Figures").mkdir(parents=True)
    SRC.mkdir()
    ref = SRC / "reference_A4.docx"
    reference_docx(ref)

    main_md = convert_main((MS / "main_q1.md").read_text(encoding="utf-8"))
    (SRC / f"PMEA_Manuscript_{V}.md").write_text(main_md, encoding="utf-8")
    to_docx(SRC / f"PMEA_Manuscript_{V}.md", OUT / f"PMEA_Manuscript_{V}.docx", ref)

    supp = (MS / "supplement_q1.md").read_text(encoding="utf-8").replace("# Supplementary Material", "# Supplementary Material 1", 1)
    title = main_md.split("\n", 1)[0].lstrip("# ").strip()
    supp = supp.replace("# Supplementary Material 1", f"# Supplementary Material 1\n\n**{title}**\n\nMehmet Kuntalp", 1)
    (SRC / f"PMEA_Supplementary_Material_1_{V}.md").write_text(supp, encoding="utf-8")
    to_docx(SRC / f"PMEA_Supplementary_Material_1_{V}.md", OUT / f"PMEA_Supplementary_Material_1_{V}.docx", ref,
            base_pt=7.5, wide_pt=6.5)

    shutil.copy(MS / "tripod_ai_checklist_q1.md", SRC / f"PMEA_TRIPOD_AI_Checklist_{V}.md")
    to_docx(SRC / f"PMEA_TRIPOD_AI_Checklist_{V}.md", OUT / f"PMEA_TRIPOD_AI_Checklist_{V}.docx", ref, base_pt=8, wide_pt=8)

    shutil.copy(MS / "PMEA_cover_letter_draft.md", SRC / f"PMEA_Cover_Letter_{V}.md")
    to_docx(SRC / f"PMEA_Cover_Letter_{V}.md", OUT / f"PMEA_Cover_Letter_{V}.docx", ref)

    (SRC / f"PMEA_Title_Page_{V}.md").write_text(title_page(title), encoding="utf-8")
    to_docx(SRC / f"PMEA_Title_Page_{V}.md", OUT / f"PMEA_Title_Page_{V}.docx", ref)
    (SRC / f"PMEA_Author_Information_and_Declarations_{V}.md").write_text(declarations_form(main_md), encoding="utf-8")
    to_docx(SRC / f"PMEA_Author_Information_and_Declarations_{V}.md", OUT / f"PMEA_Author_Information_and_Declarations_{V}.docx", ref)

    for i in (1, 2, 3):
        shutil.copy(ROOT / "figures" / f"q1_figure{i}.pdf", OUT / "PMEA_Main_Figures" / f"Figure_{i}.pdf")
    shutil.copy(MS / "PMEA_Gonderim_Kontrol_Listesi.md", OUT / f"PMEA_Gonderim_Kontrol_Listesi_{V}.md")
    previews()

    files = sorted(p for p in OUT.rglob("*") if p.is_file() and p.name != f"PMEA_{V}_manifest.json")
    manifest = {"built": date.today().isoformat(), "builder": "evaluation/build_pmea_package.py",
                "files": {str(p.relative_to(OUT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}}
    (OUT / f"PMEA_{V}_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    for p in files:
        print(f"{p.stat().st_size:>9,}  {p.relative_to(OUT)}")


if __name__ == "__main__":
    main()
